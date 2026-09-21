"""Optional GUI smoke test. Uses offscreen Qt; never talks to a sensor."""

from __future__ import annotations

import os
from pathlib import Path

import pytest
from openpyxl import load_workbook

from crfs_iq_recorder.request_builder import default_payload


@pytest.fixture
def qapp():
    os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
    from PySide6.QtWidgets import QApplication
    from crfs_iq_recorder.gui import configure_appearance

    app = QApplication.instance() or QApplication([])
    configure_appearance(app)
    yield app


@pytest.fixture(autouse=True)
def isolated_recordings(tmp_path, monkeypatch):
    monkeypatch.setattr(
        "crfs_iq_recorder.recording_history.recordings_path",
        lambda: tmp_path / "recordings.json",
    )


@pytest.fixture(autouse=True)
def isolated_settings(tmp_path, monkeypatch):
    monkeypatch.setattr(
        "crfs_iq_recorder.paths.settings_path",
        lambda: tmp_path / "settings.json",
    )


@pytest.fixture(autouse=True)
def no_explorer(monkeypatch):
    monkeypatch.setattr(
        "crfs_iq_recorder.sftp_window.QDesktopServices.openUrl",
        lambda *_args, **_kwargs: True,
    )


def _wait_listing(window, qapp, count=1, timeout=3, *, idle=False):
    import time

    deadline = time.time() + timeout
    while time.time() < deadline:
        qapp.processEvents()
        ready = window.table.topLevelItemCount() >= count
        if idle:
            ready = ready and not window._busy
        if ready:
            break
        time.sleep(0.05)
    return window.table.topLevelItemCount()


def _top_texts(window, col=1):
    return [
        window.table.topLevelItem(i).text(col)
        for i in range(window.table.topLevelItemCount())
        if window.table.topLevelItem(i) is not None
    ]


def _header_text(window, col):
    item = window.table.headerItem()
    return item.text(col) if item is not None else ""


def _group_item(window, stem: str):
    for index in range(window.table.topLevelItemCount()):
        item = window.table.topLevelItem(index)
        if item is not None and item.text(1) == stem:
            return item
    return None


def test_gui_default_payload(qapp):
    from crfs_iq_recorder.gui import MainWindow

    window = MainWindow(demo=True)
    try:
        payload, error = window._try_payload()
        assert error is None
        assert payload == default_payload()
        scan = payload["remote_recording_scans"][0]
        assert scan["duration"] == scan["rate"] == scan["capture_length"] == 0.1
        assert scan["recording_format"] == "WAVE"
        assert "7.631375" in window.size_label.text()
        assert "Estimate" in window.size_label.text()
        assert window._conn.demo is True
        assert window.model_value.text() == "R100-18"
        assert window.firmware_value.text() == "2.25-325"
        assert window.serial_value.text() == "rfeyeDEMO"
        assert window.status_value.text() == "Demo"
        assert window.link_status.text() == "Demo"
        assert window.wait_spin.value() == 5
        assert window.wait_unit.currentText() == "Seconds"
        assert window._repeat_wait_s() == 5
        from PySide6.QtWidgets import QLabel

        logos = window.findChildren(QLabel, "headerLogo")
        assert logos and not logos[0].pixmap().isNull()
    finally:
        window.close()


def test_gui_recording_time_fills_all_timing_fields(qapp):
    from crfs_iq_recorder.gui import MainWindow

    window = MainWindow(demo=True)
    try:
        window.time_edit.setText("2.5")
        payload, error = window._try_payload()
        assert error is None
        scan = payload["remote_recording_scans"][0]
        assert scan["duration"] == 2.5
        assert scan["rate"] == 2.5
        assert scan["capture_length"] == 2.5
    finally:
        window.close()


def test_gui_start_end_mhz_converts_to_hz(qapp):
    from crfs_iq_recorder.gui import MainWindow

    window = MainWindow(demo=True)
    try:
        window.mode_start.setChecked(True)
        window.field_a.setText("793.75")
        window.field_b.setText("806.25")
        payload, error = window._try_payload()
        assert error is None
        scan = payload["remote_recording_scans"][0]
        assert scan["center_frequency"] == 800_000_000
        assert scan["bandwidth"] == 12_500_000
        assert scan["recording_format"] == "WAVE"
    finally:
        window.close()


def test_sftp_window_opens_today_local_date_folder(qapp):
    from datetime import datetime

    from crfs_iq_recorder.connection_state import ConnectionState
    from crfs_iq_recorder.sftp_paths import today_remdata_directory
    from crfs_iq_recorder.sftp_window import SftpWindow

    window = SftpWindow(ConnectionState(demo=True, host="192.0.2.10"))
    try:
        expected = f"/mnt/1/remdata/{datetime.now():%Y%m%d}/"
        assert window.path_label.text() == expected
        assert window.path_label.text() == today_remdata_directory()
    finally:
        window.close()


def test_sftp_window_fills_demo_listing_without_crashing(qapp):
    import time

    from crfs_iq_recorder.connection_state import ConnectionState
    from crfs_iq_recorder.sftp_window import SftpWindow

    window = SftpWindow(ConnectionState(demo=True, host="192.0.2.10"))
    try:
        _wait_listing(window, qapp)
        assert window.table.topLevelItemCount() >= 1
        assert window.table.topLevelItem(0) is not None
        window.refresh()
        deadline = time.time() + 3
        while time.time() < deadline and window._busy:
            qapp.processEvents()
            time.sleep(0.05)
        qapp.processEvents()
        assert window.table.topLevelItemCount() >= 1
    finally:
        window.close()
        qapp.processEvents()


def test_gui_marks_recording_finished_after_recording_time(qapp):
    import time

    from crfs_iq_recorder.gui import MainWindow

    window = MainWindow(demo=True)
    try:
        window.time_edit.setText("0.1")
        window._on_start()
        assert window.start_btn.text().startswith("Recording")
        deadline = time.time() + 3
        while time.time() < deadline and window.start_btn.text() != "Start Recording":
            qapp.processEvents()
            time.sleep(0.05)
        assert window.start_btn.text() == "Start Recording"
        assert window.start_btn.isEnabled()
        text = window.msg.text().lower()
        assert "elapsed" in text
        assert "finished" not in text
        assert window._http_in_flight is False
        assert window._record_phase in {"finalizing", "unknown", "confirmed"}
        window._on_start()
        assert window.start_btn.text().startswith("Recording")
    finally:
        window.close()
        qapp.processEvents()


def test_hung_http_does_not_block_the_next_recording(qapp):
    from crfs_iq_recorder.gui import MainWindow

    window = MainWindow(demo=True)
    try:
        window._http_in_flight = True
        window._record_client = window._client()
        window._on_recording_time_elapsed()
        assert window._http_in_flight is False
        assert window.start_btn.isEnabled()
        window._on_start()
        assert window.start_btn.text().startswith("Recording")
    finally:
        window.close()
        qapp.processEvents()


def test_gui_saves_recording_parameters(qapp, tmp_path):
    from crfs_iq_recorder.gui import MainWindow
    from crfs_iq_recorder.recording_history import load_recordings

    window = MainWindow(demo=True)
    try:
        window.mode_start.setChecked(True)
        window.field_a.setText("790")
        window.field_b.setText("810")
        window.time_edit.setText("4")
        window._on_start()
        records = load_recordings(tmp_path / "recordings.json")
        assert len(records) == 1
        rec = records[0]
        assert rec.start_hz == 790_000_000
        assert rec.end_hz == 810_000_000
        assert rec.center_hz == 800_000_000
        assert rec.bandwidth_hz == 20_000_000
        assert rec.duration_s == 4
        assert rec.host == window._conn.host
    finally:
        window.close()
        qapp.processEvents()


def test_sftp_table_shows_saved_recording_parameters(qapp, tmp_path):
    import time
    from datetime import datetime

    from crfs_iq_recorder.connection_state import ConnectionState
    from crfs_iq_recorder.recording_history import add_recording, new_recording
    from crfs_iq_recorder.sftp_window import SftpWindow

    add_recording(
        new_recording(
            host="192.0.2.10",
            start_hz=790_000_000,
            end_hz=810_000_000,
            center_hz=800_000_000,
            bandwidth_hz=20_000_000,
            duration_s=4,
            started_at=datetime.now(),
            demo=True,
        )
    )
    window = SftpWindow(ConnectionState(demo=True, host="192.0.2.10"))
    try:
        _wait_listing(window, qapp)
        assert window.table.columnCount() == 9
        shown = [
            (item.text(4), item.text(8))
            for item in (
                window.table.topLevelItem(i) for i in range(window.table.topLevelItemCount())
            )
            if item is not None
        ]
        assert any(start == "790" and duration == "4" for start, duration in shown)
    finally:
        window.close()
        qapp.processEvents()


def test_sftp_downloads_multiple_files(qapp, tmp_path):
    import time

    from crfs_iq_recorder.connection_state import ConnectionState
    from crfs_iq_recorder.sftp_window import SftpWindow

    window = SftpWindow(ConnectionState(demo=True, host="192.0.2.10"))
    try:
        _wait_listing(window, qapp, 2)
        files = [entry for entry in window._entries if not entry.is_dir]
        assert len(files) >= 2
        jobs = [(entry, str(tmp_path / entry.name)) for entry in files[:2]]
        window._queue_downloads(jobs)
        deadline = time.time() + 4
        while time.time() < deadline:
            qapp.processEvents()
            if (
                (tmp_path / files[0].name).is_file()
                and (tmp_path / files[1].name).is_file()
                and not window._dl_active
            ):
                break
            time.sleep(0.05)
        assert (tmp_path / files[0].name).is_file()
        assert (tmp_path / files[1].name).is_file()
        assert window.download_btn.text() == "Download"
    finally:
        window.close()
        qapp.processEvents()


def test_sftp_begin_downloads_uses_fixed_folder_without_dialogs(qapp, tmp_path, monkeypatch):
    import time

    from PySide6.QtWidgets import QFileDialog

    from crfs_iq_recorder.connection_state import ConnectionState
    from crfs_iq_recorder.sftp_window import SftpWindow

    def _no_dialog(*_args, **_kwargs):
        raise AssertionError("Downloads must not ask for a destination.")

    monkeypatch.setattr(QFileDialog, "getSaveFileName", _no_dialog)
    monkeypatch.setattr(QFileDialog, "getExistingDirectory", _no_dialog)

    dest = tmp_path / "CRFS IQ Recorder" / "Recordings"
    window = SftpWindow(
        ConnectionState(demo=True, host="192.0.2.10", download_dir=str(dest))
    )
    try:
        _wait_listing(window, qapp, 2)
        files = [entry for entry in window._entries if not entry.is_dir]
        assert len(files) >= 2
        window._begin_downloads(files[:2])
        deadline = time.time() + 4
        while time.time() < deadline:
            qapp.processEvents()
            if (
                (dest / files[0].name).is_file()
                and (dest / files[1].name).is_file()
                and not window._dl_active
                and not window._busy
            ):
                break
            time.sleep(0.05)
        assert dest.is_dir()
        assert (dest / files[0].name).is_file()
        assert (dest / files[1].name).is_file()
        existing = dest / files[0].name
        window._begin_downloads(files[:1])
        deadline = time.time() + 4
        renamed = dest / f"{Path(files[0].name).stem} (2){Path(files[0].name).suffix}"
        while time.time() < deadline:
            qapp.processEvents()
            if renamed.is_file() and not window._dl_active and not window._busy:
                break
            time.sleep(0.05)
        assert existing.is_file()
        assert renamed.is_file()
    finally:
        window.close()
        qapp.processEvents()


def test_settings_dialog_shows_and_returns_download_folder(qapp, tmp_path):
    from crfs_iq_recorder.connection_dialog import ConnectionDialog
    from crfs_iq_recorder.connection_state import ConnectionState
    from crfs_iq_recorder.paths import default_download_dir

    dialog = ConnectionDialog(ConnectionState())
    try:
        assert dialog.download_edit.text() == str(default_download_dir())
        chosen = tmp_path / "My Recordings"
        dialog.download_edit.setText(str(chosen))
        result = dialog.result_state()
        assert result.download_dir == str(chosen)
        assert not hasattr(dialog, "demo_check")
        assert not hasattr(dialog, "sensor_combo")
        assert not hasattr(dialog, "sftp_same")
        assert dialog.sftp_user.isEnabled()
        assert dialog.sftp_pass.isEnabled()
        assert result.sftp_same_as_http is False
        assert dialog.scroll is not None
        from PySide6.QtWidgets import QDialogButtonBox

        save = dialog.buttons.button(QDialogButtonBox.StandardButton.Ok)
        assert save is not None
        assert save.text() == "Save"
        dialog.show()
        qapp.processEvents()
        dialog.resize(480, 320)
        qapp.processEvents()
        assert dialog.buttons.isVisible()
        assert dialog.scroll.isVisible()
    finally:
        dialog.close()


def test_gui_mode_switch_converts_fields(qapp):
    from crfs_iq_recorder.gui import MainWindow

    window = MainWindow(demo=True)
    try:
        window.mode_center.setChecked(True)
        window.field_a.setText("800")
        window.field_b.setText("12.5")
        window.mode_start.setChecked(True)
        assert window.field_a.text() == "793.75"
        assert window.field_b.text() == "806.25"
        payload, error = window._try_payload()
        assert error is None
        scan = payload["remote_recording_scans"][0]
        assert scan["center_frequency"] == 800_000_000
        assert scan["bandwidth"] == 12_500_000
        window.mode_center.setChecked(True)
        assert window.field_a.text() == "800"
        assert window.field_b.text() == "12.5"
    finally:
        window.close()


def test_gui_rejects_invalid_frequency(qapp):
    from crfs_iq_recorder.gui import MainWindow

    window = MainWindow(demo=True)
    try:
        window.field_a.setText("abc")
        payload, error = window._try_payload()
        assert payload is None
        assert error
        window.field_a.setText("800")
        window.field_b.setText("12.5")
        window.time_edit.setText("0")
        payload, error = window._try_payload()
        assert payload is None
    finally:
        window.close()


def test_gui_start_uses_unique_task_id_and_blocks_duplicates(qapp):
    from crfs_iq_recorder.gui import MainWindow

    window = MainWindow(demo=True)
    try:
        window.time_edit.setText("2")
        window._on_start()
        assert window._recording_active
        first_gen = window._record_gen
        window._on_start()
        assert window._record_gen == first_gen
        assert window.start_btn.isEnabled() is False
    finally:
        window.close()
        qapp.processEvents()


def test_stale_sensor_info_does_not_overwrite(qapp):
    from crfs_iq_recorder.gui import MainWindow
    from crfs_iq_recorder.sensor_info import SensorInfo

    window = MainWindow(demo=True)
    try:
        window._conn.host = "192.0.2.10"
        window._info_token = 3
        window._apply_sensor_info(
            SensorInfo(
                host="192.0.2.10",
                model="R40-8",
                firmware="2.25-325",
                serial="rfeyeTEST",
                status="Connected",
                connected=True,
            )
        )
        window._show_sensor_info_result(
            (
                2,
                "192.0.2.12",
                SensorInfo(
                    host="192.0.2.12",
                    model="StaleModel",
                    firmware="old",
                    serial="stale",
                    status="Connected",
                    connected=True,
                ),
            )
        )
        assert window.model_value.text() != "StaleModel"
        assert window.serial_value.text() == "rfeyeTEST"
        window._show_sensor_info(
            SensorInfo(
                host="192.0.2.12",
                model="WrongHost",
                firmware="x",
                serial="y",
                status="Connected",
                connected=True,
            )
        )
        assert window.model_value.text() != "WrongHost"
    finally:
        window.close()


def test_link_indicators_follow_failures_and_restore(qapp):
    from crfs_iq_recorder.gui import MainWindow

    window = MainWindow(demo=True)
    logged: list[str] = []
    original_log = window.log

    def capture(message: str, *, level: str = "info") -> None:
        logged.append(message)
        original_log(message, level=level)

    try:
        window._link_timer.stop()
        window.log = capture  # type: ignore[method-assign]
        window._refresh_sensor_info = lambda: None  # type: ignore[method-assign]
        window._refresh_storage = lambda: None  # type: ignore[method-assign]
        window._link.demo = False
        window._link.state = "connected"
        window._link.failures = 0
        window._sync_link_indicators()
        token = window._link_token
        window._apply_link_result(token, True, "ok")
        assert window.status_value.text() == "Connected"
        assert window.link_status.text() == "Connected"
        window._apply_link_result(token, False, "timeout")
        assert window.status_value.text() == "Reconnecting..."
        assert window.link_status.text() == "Reconnecting..."
        window._apply_link_result(token, False, "timeout")
        assert window.status_value.text() == "Disconnected"
        assert window.link_status.text() == "Disconnected"
        window._apply_link_result(token, False, "timeout")
        assert window.status_value.text() == "Disconnected"
        window._apply_link_result(token, True, "ok")
        assert window.status_value.text() == "Connected"
        assert window.link_status.text() == "Connected"
        assert logged.count("Sensor reconnecting...") == 1
        assert logged.count("Sensor disconnected.") == 1
        assert logged.count("Sensor connected.") == 1
        stale = token
        window._restart_link_monitor()
        window._link.demo = False
        window._link.state = "checking"
        window._sync_link_indicators()
        assert window._link_token != stale
        window._apply_link_result(stale, True, "ok")
        assert window.status_value.text() == "Checking…"
        assert window.link_status.text() == "Checking…"
        jobs = len(window._jobs)
        window._link_in_flight = True
        window._poll_link()
        assert len(window._jobs) == jobs
    finally:
        window.close()


def test_storage_failure_does_not_change_api_link(qapp):
    from crfs_iq_recorder.gui import MainWindow
    from crfs_iq_recorder.sensor_storage import StorageSnapshot

    window = MainWindow(demo=True)
    try:
        window._link_timer.stop()
        window._link.demo = False
        window._link.state = "connected"
        window._link.failures = 0
        window._sync_link_indicators()
        token = window._storage_token
        window._show_storage_result((token, StorageSnapshot.unavailable("SFTP failed")))
        assert window.status_value.text() == "Connected"
        assert window.link_status.text() == "Connected"
    finally:
        window.close()


def test_sftp_lists_newest_files_first_without_new_on_first_listing(qapp):
    import time
    from datetime import timedelta

    from crfs_iq_recorder.connection_state import ConnectionState
    from crfs_iq_recorder.sftp_client import RemoteEntry
    from crfs_iq_recorder.sftp_window import SftpWindow

    window = SftpWindow(ConnectionState(demo=True, host="192.0.2.10"))
    try:
        _wait_listing(window, qapp)
        files = [entry for entry in window._entries if not entry.is_dir]
        assert files
        names = [text.rstrip("/") for text in _top_texts(window)]
        assert names[0] == "iq-demo-2.wav"
        assert "iq_20260910_161629" in names
        new_flags = _top_texts(window, 0)
        assert all(flag == "" for flag in new_flags)
        assert "▼" in _header_text(window, 3)
        keep = window.table.topLevelItem(0)
        assert keep is not None
        window.table.clearSelection()
        keep.setSelected(True)
        window.table.setCurrentItem(keep)
        newest = max((entry.modified for entry in files if entry.modified), default=None)
        extra = RemoteEntry(
            "iq_brand_new.wav",
            files[0].path.rsplit("/", 1)[0] + "/iq_brand_new.wav",
            False,
            100,
            newest + timedelta(seconds=30) if newest else None,
        )
        window._fill(window._entries + [extra])
        assert window.table.topLevelItem(0).text(1) == "iq_brand_new.wav"
        assert window.table.topLevelItem(0).text(0) == "NEW"
        selected = window._selected_file_paths()
        assert any(path.endswith("iq-demo-2.wav") for path in selected)
        assert extra.path not in selected
    finally:
        window.close()
        qapp.processEvents()


def test_gui_ghz_unit_converts_to_hz(qapp):
    from crfs_iq_recorder.gui import MainWindow

    window = MainWindow(demo=True)
    try:
        window.mode_center.setChecked(True)
        window.unit_combo.setCurrentText("GHz")
        window.field_a.setText("0.8")
        window.field_b.setText("0.0125")
        payload, error = window._try_payload()
        assert error is None
        scan = payload["remote_recording_scans"][0]
        assert scan["center_frequency"] == 800_000_000
        assert scan["bandwidth"] == 12_500_000
        assert window.unit_a.text() == "GHz"
        assert "Center frequency" in window.label_a.text()
    finally:
        window.close()


def test_gui_verified_format_is_sent(qapp):
    from crfs_iq_recorder.gui import MainWindow

    window = MainWindow(demo=True)
    try:
        window.format_combo.setCurrentText("NCP")
        payload, error = window._try_payload()
        assert error is None
        assert payload["remote_recording_scans"][0]["recording_format"] == "NCP"
        labels = [window.format_combo.itemText(i) for i in range(window.format_combo.count())]
        assert labels == ["WAVE", "XDAT", "NCP", "HDF5"]
    finally:
        window.close()


def test_gui_remembers_band_time_unit_format(qapp):
    from crfs_iq_recorder.gui import MainWindow

    window = MainWindow(demo=True)
    try:
        window.mode_start.setChecked(True)
        window.unit_combo.setCurrentText("GHz")
        window.field_a.setText("0.79")
        window.field_b.setText("0.81")
        window.time_edit.setText("4")
        window.format_combo.setCurrentText("XDAT")
    finally:
        window.close()
        qapp.processEvents()

    window2 = MainWindow(demo=True)
    try:
        assert window2.mode_start.isChecked()
        assert window2.unit_combo.currentText() == "GHz"
        assert window2.time_edit.text() == "4"
        assert window2.format_combo.currentText() == "XDAT"
        payload, error = window2._try_payload()
        assert error is None
        scan = payload["remote_recording_scans"][0]
        assert scan["center_frequency"] == 800_000_000
        assert scan["bandwidth"] == 20_000_000
        assert scan["recording_format"] == "XDAT"
        assert scan["duration"] == 4
    finally:
        window2.close()


def test_download_parent_downloads_split_parts(qapp):
    import time

    from crfs_iq_recorder.connection_state import ConnectionState
    from crfs_iq_recorder.sftp_client import RemoteEntry
    from crfs_iq_recorder.sftp_window import SftpWindow

    window = SftpWindow(ConnectionState(demo=True, host="192.0.2.10"))
    try:
        _wait_listing(window, qapp)
        group = _group_item(window, "iq_20260910_161629")
        assert group is not None
        assert group.childCount() == 2
        assert group.child(0).text(1) == "_0001.wav"
        assert group.child(1).text(1) == "_0002.wav"
        window.table.clearSelection()
        group.setSelected(True)
        window.table.setCurrentItem(group)
        window._on_selection_changed()
        assert window.download_btn.text() == "Download this recording"
        captured = []
        window._begin_downloads = lambda files: captured.extend(files)
        window._download_selected()
        names = {entry.name for entry in captured}
        assert names == {"iq_20260910_161629_0001.wav", "iq_20260910_161629_0002.wav"}
        captured.clear()
        window.table.clearSelection()
        group.setExpanded(True)
        child = group.child(0)
        child.setSelected(True)
        window.table.setCurrentItem(child)
        window._on_selection_changed()
        assert window.download_btn.text() == "Download"
        window._download_selected()
        assert [entry.name for entry in captured] == ["iq_20260910_161629_0001.wav"]
        assert window.open_folder_btn.text() == "Open recordings folder"
        window.table.clearSelection()
        group.setExpanded(True)
        group.setSelected(True)
        group.child(0).setSelected(True)
        window._on_selection_changed()
        assert set(window._selected_file_paths()) == {
            next(entry.path for entry in window._entries if entry.name.endswith("161629_0001.wav")),
            next(entry.path for entry in window._entries if entry.name.endswith("161629_0002.wav")),
        }
        group.setExpanded(True)
        qapp.processEvents()
        part1 = next(entry for entry in window._entries if entry.name.endswith("161629_0001.wav"))
        extra = RemoteEntry(
            "iq_20260910_161629_0003.wav",
            part1.path.rsplit("/", 1)[0] + "/iq_20260910_161629_0003.wav",
            False,
            100,
            part1.modified,
        )
        window._fill(window._entries + [extra])
        group = _group_item(window, "iq_20260910_161629")
        assert group is not None
        assert group.childCount() == 3
        assert [group.child(i).text(1) for i in range(3)] == ["_0001.wav", "_0002.wav", "_0003.wav"]
        assert group.isExpanded()
    finally:
        window.close()
        qapp.processEvents()


def test_open_recordings_folder_uses_saved_download_dir(qapp, tmp_path, monkeypatch):
    opened = []

    def _open(url):
        opened.append(url.toLocalFile() if hasattr(url, "toLocalFile") else str(url))
        return True

    monkeypatch.setattr("crfs_iq_recorder.sftp_window.QDesktopServices.openUrl", _open)
    dest = tmp_path / "CRFS IQ Recorder" / "Recordings"
    dest.mkdir(parents=True)
    from crfs_iq_recorder.connection_state import ConnectionState
    from crfs_iq_recorder.sftp_window import SftpWindow

    window = SftpWindow(ConnectionState(demo=True, host="192.0.2.10", download_dir=str(dest)))
    try:
        window._open_recordings_folder()
        assert opened
        assert Path(opened[0]).resolve() == dest.resolve()
    finally:
        window.close()
        qapp.processEvents()


def _luminance(color) -> float:
    return 0.299 * color.red() + 0.587 * color.green() + 0.114 * color.blue()


def test_combo_popups_have_readable_contrast(qapp):
    from PySide6.QtGui import QPalette

    from crfs_iq_recorder.gui import MainWindow

    window = MainWindow(demo=True)
    try:
        form = window.format_combo.view().palette()
        form_bg = form.color(QPalette.ColorRole.Base)
        form_fg = form.color(QPalette.ColorRole.Text)
        assert abs(_luminance(form_bg) - _luminance(form_fg)) > 80
        assert _luminance(form_bg) > 180
        assert _luminance(form_fg) < 80
        units = [window.unit_combo.itemText(i) for i in range(window.unit_combo.count())]
        assert units == ["kHz", "MHz", "GHz"]
        assert not hasattr(window, "sensor_combo")
    finally:
        window.close()


def test_recording_fields_do_not_overlap(qapp):
    from PySide6.QtCore import QRect

    from crfs_iq_recorder.gui import MainWindow

    window = MainWindow(demo=True)
    try:
        for size in ((1100, 720), (880, 560)):
            window.resize(*size)
            window.show()
            qapp.processEvents()
            widgets = [window.field_a, window.field_b, window.format_combo, window.time_edit]
            rects = []
            for widget in widgets:
                top = widget.mapTo(window, widget.rect().topLeft())
                rects.append(QRect(top.x(), top.y(), widget.width(), widget.height()))
            for index, first in enumerate(rects):
                for second in rects[index + 1 :]:
                    assert not first.intersects(second), (size, first, second)
            assert window.field_a.height() >= 28
            assert "800" in window.field_a.text() or window.field_a.text()
            assert window.format_combo.currentText() == "WAVE"
    finally:
        window.close()


def test_sftp_up_opens_parent_folder(qapp):
    import time

    from crfs_iq_recorder.connection_state import ConnectionState
    from crfs_iq_recorder.sftp_paths import parent_directory, today_remdata_directory
    from crfs_iq_recorder.sftp_window import SftpWindow

    window = SftpWindow(ConnectionState(demo=True, host="192.0.2.10"))
    try:
        _wait_listing(window, qapp)
        today = today_remdata_directory()
        assert window.path_label.text().rstrip("/") == today.rstrip("/")

        def wait_path(expected: str) -> bool:
            want = expected.rstrip("/") or "/"
            deadline = time.time() + 3
            while time.time() < deadline:
                qapp.processEvents()
                current = window.path_label.text().rstrip("/") or "/"
                if current == want and not window._busy:
                    return True
                time.sleep(0.05)
            return False

        window._go_parent()
        parent = parent_directory(today)
        assert wait_path(parent)
        names = [text.rstrip("/") for text in _top_texts(window)]
        assert any(name.isdigit() and len(name) == 8 for name in names)
        today_name = today.rstrip("/").rsplit("/", 1)[-1]
        assert names[0] == today_name
        folder_dates = [name for name in names if name.isdigit() and len(name) == 8]
        assert folder_dates == sorted(folder_dates, reverse=True)
        for expected in ("/mnt/1/", "/mnt/", "/"):
            window._go_parent()
            assert wait_path(expected)
        window._go_parent()
        qapp.processEvents()
        assert (window.path_label.text().rstrip("/") or "/") == "/"
    finally:
        window.close()
        qapp.processEvents()


def test_theme_toggle_persists(qapp, tmp_path, monkeypatch):
    from crfs_iq_recorder.gui import MainWindow
    from crfs_iq_recorder.paths import settings_path
    from crfs_iq_recorder.theme import DARK_MODE, LIGHT_MODE, current_mode

    window = MainWindow(demo=True)
    try:
        assert current_mode() == LIGHT_MODE
        window._toggle_theme()
        assert current_mode() == DARK_MODE
        assert window.theme_btn.text() == ""
        assert window.theme_btn.toolTip() == "Switch to Light Mode"
        assert not window.theme_btn.icon().isNull()
        saved = settings_path().read_text(encoding="utf-8")
        assert '"ui_theme": "dark"' in saved
        window._toggle_theme()
        assert current_mode() == LIGHT_MODE
        assert window.theme_btn.toolTip() == "Switch to Dark Mode"
    finally:
        from crfs_iq_recorder.gui import configure_appearance

        configure_appearance(qapp, LIGHT_MODE)
        window.close()


def test_sftp_column_sort_persists_until_directory_change(qapp):
    import time

    from crfs_iq_recorder.connection_state import ConnectionState
    from crfs_iq_recorder.listing_sort import COL_NAME, COL_SIZE
    from crfs_iq_recorder.sftp_paths import parent_directory, today_remdata_directory
    from crfs_iq_recorder.sftp_window import SftpWindow

    window = SftpWindow(ConnectionState(demo=True, host="192.0.2.10"))
    try:
        _wait_listing(window, qapp, idle=True)
        assert window.table.topLevelItemCount() > 0
        assert window._user_sort is False
        window._on_header_clicked(COL_NAME)
        qapp.processEvents()
        names = [text.rstrip("/") for text in _top_texts(window)]
        assert names == sorted(names, key=str.casefold)
        assert "▲" in _header_text(window, COL_NAME)
        window._on_header_clicked(COL_NAME)
        qapp.processEvents()
        names_desc = [text.rstrip("/") for text in _top_texts(window)]
        assert names_desc == sorted(names, key=str.casefold, reverse=True)
        window._on_header_clicked(COL_SIZE)
        qapp.processEvents()
        assert window._user_sort is True
        assert window._sort_column == COL_SIZE
        snapshot = list(window._entries)
        window._fill(snapshot)
        assert window._user_sort is True
        assert window._sort_column == COL_SIZE
        window._go_parent()
        expected = parent_directory(today_remdata_directory())
        deadline = time.time() + 3
        while time.time() < deadline:
            qapp.processEvents()
            current = window.path_label.text().rstrip("/") or "/"
            if current == expected.rstrip("/") and not window._busy:
                break
            time.sleep(0.05)
        assert window._user_sort is False
        assert "▼" in _header_text(window, 3)
    finally:
        window.close()
        qapp.processEvents()


def _drive_capture(window, qapp) -> None:
    qapp.processEvents()
    window._record_timer.stop()
    if window._recording_active:
        window._on_recording_time_elapsed()
    qapp.processEvents()
    for _ in range(4):
        window._poll_collection_files()
    qapp.processEvents()


def test_gui_suggests_iridium_and_logs_repeat_parts(qapp, tmp_path):
    from crfs_iq_recorder.gui import MainWindow
    from test_collection_excel import build_sample_workbook

    path = build_sample_workbook(tmp_path / "collection.xlsm")
    window = MainWindow(demo=True)
    try:
        window._settle_s = 0
        window._stable_needed = 2
        window._demo_parts = 3
        window.excel_check.setChecked(True)
        assert window._load_workbook(str(path), quiet=True)
        window.mode_start.setChecked(True)
        window.unit_combo.setCurrentText("MHz")
        window.field_a.setText("1616")
        window.field_b.setText("1626")
        window._refresh_class_suggestion()
        assert window.class_combo.currentText() == "Iridium"
        assert window.sheet_combo.currentText() == "SATCOM"
        window.time_edit.setText("0.05")
        window.wait_spin.setValue(0)
        window.repeat_n.setChecked(True)
        window.repeat_spin.setValue(2)
        window._on_start()
        assert window._series_active
        _drive_capture(window, qapp)
        _drive_capture(window, qapp)
        assert window._recordings_completed == 2
        assert window._files_found == 6
        assert window._series_active is False
        from crfs_iq_recorder.recording_history import split_iq_filenames

        wb = load_workbook(path)
        iq_cells = [
            cell
            for cell in wb["SATCOM"]["I"]
            if cell.row > 1 and any(name.lower().endswith(".wav") for name in split_iq_filenames(cell.value))
        ]
        assert len(iq_cells) == 2
        all_names = [name for cell in iq_cells for name in split_iq_filenames(cell.value)]
        assert len(all_names) == 6
        assert len(set(all_names)) == 6
        assert all(str(cell.value).count("\n") == 2 for cell in iq_cells)
        logged_classes = [wb["SATCOM"][f"B{cell.row}"].value for cell in iq_cells]
        assert set(logged_classes) == {"Iridium"}
        starts = [wb["SATCOM"][f"C{cell.row}"].value for cell in iq_cells]
        assert starts[-1] == "1616 MHz"
        durations = [wb["SATCOM"][f"K{cell.row}"].value for cell in iq_cells]
        assert all(value == "0.05 seconds" for value in durations)
        assert all(wb["SATCOM"][f"M{cell.row}"].value in (None, "") for cell in iq_cells)
        assert all(cell.alignment.wrap_text is True for cell in iq_cells)
        wb.close()
    finally:
        window._watch_timer.stop()
        window._record_timer.stop()
        window.close()
        qapp.processEvents()


def test_gui_suggests_satcom_for_full_range(qapp, tmp_path):
    from crfs_iq_recorder.gui import MainWindow
    from test_collection_excel import build_sample_workbook

    path = build_sample_workbook(tmp_path / "collection.xlsm")
    window = MainWindow(demo=True)
    try:
        window.excel_check.setChecked(True)
        assert window._load_workbook(str(path), quiet=True)
        window.mode_start.setChecked(True)
        window.unit_combo.setCurrentText("MHz")
        window.field_a.setText("1616")
        window.field_b.setText("1660.5")
        window._refresh_class_suggestion()
        assert window.class_combo.currentText() == "Satcom"
        assert window.sheet_combo.currentText() == "SATCOM"
    finally:
        window.close()
        qapp.processEvents()


def test_gui_stop_after_current_does_not_start_next(qapp, tmp_path):
    from crfs_iq_recorder.gui import MainWindow
    from test_collection_excel import build_sample_workbook

    path = build_sample_workbook(tmp_path / "collection.xlsm")
    window = MainWindow(demo=True)
    try:
        window._settle_s = 0
        window._stable_needed = 2
        window._demo_parts = 1
        window.excel_check.setChecked(True)
        window._load_workbook(str(path), quiet=True)
        window.mode_start.setChecked(True)
        window.field_a.setText("156")
        window.field_b.setText("162")
        window.time_edit.setText("0.05")
        window.wait_spin.setValue(0)
        window.repeat_until.setChecked(True)
        window._on_start()
        window._on_stop_after_current()
        assert window._stop_after_current
        _drive_capture(window, qapp)
        qapp.processEvents()
        assert window._recordings_completed == 1
        assert window._series_active is False
        assert window.start_btn.isEnabled()
    finally:
        window._watch_timer.stop()
        window._record_timer.stop()
        window.close()
        qapp.processEvents()


def test_gui_wait_minutes_and_stop_during_countdown(qapp, tmp_path):
    from crfs_iq_recorder.gui import MainWindow
    from test_collection_excel import build_sample_workbook

    path = build_sample_workbook(tmp_path / "collection.xlsm")
    window = MainWindow(demo=True)
    try:
        window._settle_s = 0
        window._stable_needed = 2
        window._demo_parts = 1
        window.excel_check.setChecked(True)
        window._load_workbook(str(path), quiet=True)
        window.mode_start.setChecked(True)
        window.field_a.setText("156")
        window.field_b.setText("162")
        window.time_edit.setText("0.05")
        window.wait_unit.setCurrentText("Minutes")
        window.wait_spin.setValue(2)
        assert window._repeat_wait_s() == 120
        window.wait_unit.setCurrentText("Seconds")
        window.wait_spin.setValue(30)
        assert window._repeat_wait_s() == 30
        window.repeat_until.setChecked(True)
        window._on_start()
        _drive_capture(window, qapp)
        assert window._recordings_completed == 1
        assert window._waiting is True
        assert window._series_active is True
        assert "Next recording in" in (window.phase_label.text() + window.start_btn.text())
        window._on_stop_after_current()
        qapp.processEvents()
        assert window._waiting is False
        assert window._series_active is False
        assert window._recordings_completed == 1
    finally:
        window._wait_timer.stop()
        window._watch_timer.stop()
        window._record_timer.stop()
        window.close()
        qapp.processEvents()


def test_excel_section_collapses_without_losing_values(qapp, tmp_path):
    from PySide6.QtCore import QPoint

    from crfs_iq_recorder.gui import MainWindow
    from test_collection_excel import build_sample_workbook

    path = build_sample_workbook(tmp_path / "collection.xlsm")
    window = MainWindow(demo=True)
    try:
        window.show()
        qapp.processEvents()
        assert window.excel_check.isHidden() is False
        assert window.excel_details.isHidden()
        window.excel_check.setChecked(True)
        assert window._load_workbook(str(path), quiet=True)
        window.event_combo.setEditText("Haifa Port")
        window.class_combo.setEditText("Iridium")
        window.repeat_n.setChecked(True)
        window.repeat_spin.setValue(7)
        window.wait_spin.setValue(12)
        window.wait_unit.setCurrentText("Seconds")
        qapp.processEvents()
        assert not window.excel_details.isHidden()
        expanded = window.excel_section.height()
        start_y = window.start_btn.mapTo(window, QPoint(0, 0)).y()
        window.excel_check.setChecked(False)
        qapp.processEvents()
        assert window.excel_check.isHidden() is False
        assert window.excel_details.isHidden()
        assert window.excel_section.height() < expanded
        assert window.start_btn.mapTo(window, QPoint(0, 0)).y() < start_y
        window.excel_check.setChecked(True)
        qapp.processEvents()
        assert not window.excel_details.isHidden()
        assert window.workbook_edit.text() == str(path)
        assert window.event_combo.currentText() == "Haifa Port"
        assert window.class_combo.currentText() == "Iridium"
        assert window.repeat_n.isChecked()
        assert window.repeat_spin.value() == 7
        assert window.wait_spin.value() == 12
    finally:
        window.close()
        qapp.processEvents()


def test_gui_remembers_wait_and_credentials(qapp):
    from crfs_iq_recorder.gui import MainWindow

    window = MainWindow(demo=True)
    try:
        window.wait_unit.setCurrentText("Minutes")
        window.wait_spin.setValue(7)
        window._conn.http_password = "custom-http"
        window._conn.sftp_password = "custom-sftp"
    finally:
        window.close()
        qapp.processEvents()
    restored = MainWindow(demo=True)
    try:
        assert restored.wait_unit.currentText() == "Minutes"
        assert restored._repeat_wait_s() == 420
        assert restored._conn.http_password == "custom-http"
        assert restored._conn.sftp_password == "custom-sftp"
        assert restored.wait_spin.value() == 7
    finally:
        restored.close()
        qapp.processEvents()


def test_sftp_delete_removes_selected_files_only(qapp, monkeypatch):
    import time

    from crfs_iq_recorder.connection_state import ConnectionState
    from crfs_iq_recorder.demo_sftp import shared_demo_browser
    from crfs_iq_recorder.sftp_client import join_remote
    from crfs_iq_recorder.sftp_paths import today_remdata_directory
    from crfs_iq_recorder.sftp_window import SftpWindow

    browser = shared_demo_browser()
    today = today_remdata_directory()
    target = join_remote(today, "iq-demo.wav")
    other = join_remote(today, "iq-demo-2.wav")
    keep = browser.publish_capture("iq_keep_local", parts=1)[0]
    window = SftpWindow(ConnectionState(demo=True, host="192.0.2.10"))
    logs: list[str] = []
    window.activity.connect(logs.append)
    monkeypatch.setattr(window.table, "hasFocus", lambda: True)
    try:
        _wait_listing(window, qapp, 3, timeout=4)
        assert window.table.topLevelItemCount() >= 3
        window.select_paths([target])
        assert window._selected_file_paths() == [target]
        window._delete_selected_remote()
        deadline = time.time() + 4
        while time.time() < deadline and target in browser._files:
            qapp.processEvents()
            time.sleep(0.05)
        assert target not in browser._files
        assert other in browser._files
        assert keep.path in browser._files
        deadline = time.time() + 4
        while time.time() < deadline:
            qapp.processEvents()
            names = [entry.name for entry in window._entries if not entry.is_dir]
            if "iq-demo.wav" not in names and "iq-demo-2.wav" in names:
                break
            time.sleep(0.05)
        names = [entry.name for entry in window._entries if not entry.is_dir]
        assert "iq-demo.wav" not in names
        assert "iq-demo-2.wav" in names
        window.select_paths([other, keep.path])
        assert set(window._selected_file_paths()) == {other, keep.path}
        window._delete_selected_remote()
        deadline = time.time() + 4
        while time.time() < deadline and (other in browser._files or keep.path in browser._files):
            qapp.processEvents()
            time.sleep(0.05)
        assert other not in browser._files
        assert keep.path not in browser._files
        assert any("Deleted" in item for item in logs)
    finally:
        window.close()
        qapp.processEvents()


def test_sftp_delete_parent_removes_all_listed_parts(qapp, monkeypatch):
    import time

    from crfs_iq_recorder.connection_state import ConnectionState
    from crfs_iq_recorder.demo_sftp import shared_demo_browser
    from crfs_iq_recorder.sftp_window import SftpWindow

    parts = shared_demo_browser().publish_capture("iq_group_delete", parts=2)
    window = SftpWindow(ConnectionState(demo=True, host="192.0.2.10"))
    monkeypatch.setattr(window.table, "hasFocus", lambda: True)
    try:
        _wait_listing(window, qapp, 1, timeout=4)
        group = _group_item(window, "iq_group_delete")
        assert group is not None
        window.table.clearSelection()
        group.setSelected(True)
        window.table.setCurrentItem(group)
        assert set(window._selected_file_paths()) == {parts[0].path, parts[1].path}
        window._delete_selected_remote()
        deadline = time.time() + 4
        browser = shared_demo_browser()
        while time.time() < deadline and any(part.path in browser._files for part in parts):
            qapp.processEvents()
            time.sleep(0.05)
        assert parts[0].path not in browser._files
        assert parts[1].path not in browser._files
    finally:
        window.close()
        qapp.processEvents()
