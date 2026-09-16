"""Class matching, frequency text, Excel append, and split-file logging."""

from __future__ import annotations

from datetime import datetime, timedelta
from pathlib import Path
from shutil import copy2

import pytest
from openpyxl import Workbook, load_workbook

from crfs_iq_recorder.class_match import suggest_class
from crfs_iq_recorder.collection_catalog import ClassMapping, load_catalog_from_rows, load_workbook_catalog
from crfs_iq_recorder.excel_log import ExcelLogStore, flush_excel_queue, pending_from_row
from crfs_iq_recorder.excel_writer import ExcelBusyError, ExcelRow, append_rows_openpyxl
from crfs_iq_recorder.file_watch import FileWatchState, capture_is_settled, update_file_watch
from crfs_iq_recorder.frequency import FreqTextStyle, format_freq_text, parse_freq_text
from crfs_iq_recorder.recording_history import split_iq_filenames
from crfs_iq_recorder.sensor_label import workbook_sensor_name
from crfs_iq_recorder.sftp_client import RemoteEntry

_HAIFA_CANDIDATES = (
    Path(r"C:\Users\itay.h\Downloads\Haifa Master 16.09.2026(1).xlsm"),
    Path(r"C:\Users\itay.h\Downloads\Haifa Master 16.09.2026.xlsm"),
)
HAIFA = next((path for path in _HAIFA_CANDIDATES if path.is_file()), _HAIFA_CANDIDATES[0])

HEADERS = [
    "Collection Event",
    "Class",
    "Freq Start",
    "Freq Stop",
    "RBW",
    "Scan Rate",
    "Sensor",
    "Snapshot ID",
    "IQ Recording",
    "IQ Image",
    "Playback Start",
    "Playback End",
    "Playback Time (UTC-3)",
]


def _mapping(sheet, name, start_hz, end_hz, start="156 MHz", stop="162 MHz") -> ClassMapping:
    return ClassMapping(
        worksheet=sheet,
        target_class=name,
        start_hz=start_hz,
        end_hz=end_hz,
        start_style=freq_style(start),
        stop_style=freq_style(stop),
    )


def freq_style(text: str) -> FreqTextStyle:
    from crfs_iq_recorder.frequency import freq_text_style

    style = freq_text_style(text)
    assert style is not None
    return style


def haifa_mappings() -> tuple[ClassMapping, ...]:
    return (
        _mapping("Tactical Radio", "Tactical Radio", 156_000_000, 162_000_000, "156 MHz", "162 MHz"),
        _mapping("SATCOM", "Iridium", 1_616_000_000, 1_626_000_000, "1616 MHz", "1626MHz"),
        _mapping("SATCOM", "Inmarsat", 1_626_500_000, 1_660_500_000, "1626.5 MHz", "1660.5MHz"),
        _mapping("SATCOM", "Satcom", 1_616_000_000, 1_660_500_000, "1616 MHz", "1660.5MHz"),
        _mapping("S-Band", "S-Band", 2_900_000_000, 3_100_000_000, "2.9 GHz", "3.1 GHz"),
        _mapping("X-Band", "X-Band", 9_200_000_000, 9_500_000_000, "9.2 GHz", "9.5 GHz"),
    )


def build_sample_workbook(path: Path) -> Path:
    wb = Workbook()
    sheets = {
        "Tactical Radio": [("Haifa Port", "Tactical Radio", "156 MHz", "162 MHz", "1 kHz", "10 Hz", "CRFS 100-18")],
        "SATCOM": [
            ("Haifa Port", "Iridium", "1616 MHz", "1626MHz", "1 kHz", "10 Hz", "CRFS 100-18"),
            ("Haifa Port", "Inmarsat", "1626.5 MHz", "1660.5MHz", "1 kHz", "10 Hz", "CRFS 100-18"),
            ("Haifa Port", "Satcom", "1616 MHz", "1660.5MHz", "", "10 Hz", "CRFS 100-18"),
            ("Haifa Port", "Satcom", "1616 MHz", "1660.5MHz", "", "", "CRFS 100-18"),
        ],
        "S-Band": [("Haifa Port", "S-Band", "2.9 GHz", "3.1 GHz", "30 kHz", "100 Hz", "CRFS 100-18")],
        "X-Band": [("Haifa Port", "X-Band", "9.2 GHz", "9.5 GHz", "30 kHz", "10 Hz", "CRFS 100-18")],
    }
    first = True
    for name, rows in sheets.items():
        ws = wb.active if first else wb.create_sheet(name)
        if first:
            ws.title = name
            first = False
        for col, header in enumerate(HEADERS, start=1):
            ws.cell(1, col, header)
        for r, row in enumerate(rows, start=2):
            for c, value in enumerate(row, start=1):
                ws.cell(r, c, value)
            if name == "Tactical Radio" and r == 2:
                ws.cell(
                    r,
                    13,
                    '=IF(K2="","",TEXT(MOD(K2-TIME(3,0,0),1),"hh:mm:ss")&IF(L2="",""," - "&TEXT(MOD(L2-TIME(3,0,0),1),"hh:mm:ss")))',
                )
        ws.auto_filter.ref = f"A1:M{1 + len(rows)}"
        ws.freeze_panes = "C2"
    wb.save(path)
    return path


def test_playback_start_is_plain_seconds_text():
    from crfs_iq_recorder.excel_writer import format_playback_duration

    assert format_playback_duration(2) == "2 seconds"
    assert format_playback_duration(4.0) == "4 seconds"
    assert format_playback_duration(0.5) == "0.5 seconds"
    assert format_playback_duration(None) is None


def test_wav_duration_from_header():
    from crfs_iq_recorder.iq_wav import stereo_iq_tone_wav_bytes, wav_duration_seconds

    payload = stereo_iq_tone_wav_bytes(seconds=0.08)
    duration = wav_duration_seconds(payload)
    assert duration is not None
    assert abs(duration - 0.08) < 0.002
    assert wav_duration_seconds(payload[:512]) == duration
    assert wav_duration_seconds(b"not-a-wav") is None


def test_parse_freq_text_allows_missing_space():
    assert parse_freq_text("1626MHz") == 1_626_000_000
    assert parse_freq_text("1626.5 MHz") == 1_626_500_000
    assert parse_freq_text("2.9 GHz") == 2_900_000_000
    assert format_freq_text(1_626_000_000, FreqTextStyle("MHz", False)) == "1626MHz"
    assert format_freq_text(1_626_500_000, FreqTextStyle("MHz", True)) == "1626.5 MHz"


def test_catalog_ignores_completed_iq_rows_as_mappings():
    header = HEADERS
    rows = [
        ["Haifa Port", "Iridium", "1616 MHz", "1626MHz", "1 kHz", "10 Hz", "CRFS 100-18", "", ""],
        [
            "Haifa Port",
            "Iridium",
            "1618 MHz",
            "1624MHz",
            "",
            "",
            "CRFS 100-18",
            "",
            "iq_logged_0001.wav",
        ],
    ]
    catalog = load_catalog_from_rows("book.xlsm", [("SATCOM", header, rows)])
    assert [item.target_class for item in catalog.mappings] == ["Iridium"]
    assert catalog.mappings[0].start_hz == 1_616_000_000
    assert catalog.mappings[0].end_hz == 1_626_000_000


def test_catalog_deduplicates_template_rows():
    header = HEADERS
    rows = [
        ["Haifa Port", "Satcom", "1616 MHz", "1660.5MHz", "", "", "CRFS 100-18"],
        ["Haifa Port", "Satcom", "1616 MHz", "1660.5MHz", "", "", "CRFS 100-18"],
        ["Haifa Port", "Iridium", "1616 MHz", "1626MHz", "1 kHz", "10 Hz", "CRFS 100-18"],
    ]
    catalog = load_catalog_from_rows("book.xlsm", [("SATCOM", header, rows)])
    names = [item.target_class for item in catalog.mappings]
    assert names == ["Satcom", "Iridium"]
    assert catalog.events == ("Haifa Port",)


def test_satcom_exact_and_narrowest_matches():
    mappings = haifa_mappings()
    wide = suggest_class(1_616_000_000, 1_660_500_000, mappings)
    assert wide.unique and wide.mapping is not None
    assert wide.mapping.target_class == "Satcom"
    iridium = suggest_class(1_616_000_000, 1_626_000_000, mappings)
    assert iridium.mapping is not None
    assert iridium.mapping.target_class == "Iridium"
    sub = suggest_class(1_618_000_000, 1_624_000_000, mappings)
    assert sub.mapping is not None
    assert sub.mapping.target_class == "Iridium"
    user = suggest_class(
        1_616_000_000,
        1_626_000_000,
        mappings,
        selected_class="Satcom",
        selected_sheet="SATCOM",
    )
    assert user.mapping is not None
    assert user.mapping.target_class == "Satcom"


def test_ambiguous_range_needs_choice():
    mappings = haifa_mappings()
    miss = suggest_class(100_000_000, 110_000_000, mappings)
    assert miss.needs_choice
    assert miss.mapping is None


def test_sensor_label_matches_workbook():
    assert workbook_sensor_name("R100-18", ["CRFS 100-18"]) == "CRFS 100-18"
    assert workbook_sensor_name("Node 100-18") == "CRFS 100-18"


def test_append_does_not_overwrite_empty_iq_template(tmp_path):
    path = build_sample_workbook(tmp_path / "sample.xlsx")
    before = load_workbook(path)
    satcom = before["SATCOM"]
    assert satcom["I2"].value in (None, "")
    assert satcom["B2"].value == "Iridium"
    last_before = satcom.max_row
    before.close()
    row = ExcelRow(
        worksheet="SATCOM",
        collection_event="Haifa Port",
        target_class="Iridium",
        start_hz=1_618_000_000,
        end_hz=1_624_000_000,
        start_style=FreqTextStyle("MHz", True),
        stop_style=FreqTextStyle("MHz", False),
        sensor="CRFS 100-18",
        filename="iq_demo_0001.wav",
        remote_path="/mnt/1/remdata/20260915/iq_demo_0001.wav",
        sensor_id="rfeyeDEMO",
        duration_s=0.5,
    )
    written = append_rows_openpyxl(path, [row])
    after = load_workbook(path)
    ws = after["SATCOM"]
    assert ws["B2"].value == "Iridium"
    assert ws["I2"].value in (None, "")
    dest = written[0]
    assert dest == last_before + 1
    assert ws.cell(dest, 9).value == "iq_demo_0001.wav"
    assert ws.cell(dest, 2).value == "Iridium"
    assert ws.cell(dest, 3).value == "1618 MHz"
    assert ws.cell(dest, 4).value == "1624MHz"
    assert ws.cell(dest, 5).value in (None, "")
    assert ws.cell(dest, 6).value in (None, "")
    assert ws.cell(dest, 11).value == "0.5 seconds"
    assert ws.cell(dest, 11).comment is None
    assert ws.cell(dest, 12).value in (None, "")
    assert ws.cell(dest, 13).value in (None, "")
    assert ws.auto_filter.ref.endswith(str(dest))
    after.close()


def test_split_parts_log_once_and_queue_survives_lock(tmp_path, monkeypatch):
    path = build_sample_workbook(tmp_path / "sample.xlsx")
    store = ExcelLogStore(tmp_path / "queue.json")
    mapping = haifa_mappings()[1]
    files = [
        ("iq_task_0001.wav", "/mnt/1/remdata/20260915/iq_task_0001.wav"),
        ("iq_task_0002.wav", "/mnt/1/remdata/20260915/iq_task_0002.wav"),
        ("iq_task_0003.wav", "/mnt/1/remdata/20260915/iq_task_0003.wav"),
    ]
    for name, remote in files:
        row = ExcelRow(
            worksheet="SATCOM",
            collection_event="Haifa Port",
            target_class="Iridium",
            start_hz=mapping.start_hz,
            end_hz=mapping.end_hz,
            start_style=mapping.start_style,
            stop_style=mapping.stop_style,
            sensor="CRFS 100-18",
            filename=name,
            remote_path=remote,
            sensor_id="rfeyeDEMO",
            duration_s=0.08,
        )
        store.enqueue(pending_from_row(workbook=path, row=row))
        store.enqueue(pending_from_row(workbook=path, row=row))
    assert store.pending_count() == 3
    saved, pending = flush_excel_queue(store)
    assert saved == 3
    assert pending == 0
    saved, pending = flush_excel_queue(store)
    assert saved == 0
    assert pending == 0
    wb = load_workbook(path)
    iq_cells = [
        cell
        for cell in wb["SATCOM"]["I"]
        if cell.row > 1 and any(name.lower().endswith(".wav") for name in split_iq_filenames(cell.value))
    ]
    assert len(iq_cells) == 1
    names = split_iq_filenames(iq_cells[0].value)
    assert names == ["iq_task_0001.wav", "iq_task_0002.wav", "iq_task_0003.wav"]
    row_idx = iq_cells[0].row
    assert iq_cells[0].alignment.wrap_text is True
    assert (wb["SATCOM"].row_dimensions[row_idx].height or 0) >= 16 * 3
    assert wb["SATCOM"].cell(row_idx, 11).value == "0.08 seconds"
    assert wb["SATCOM"].cell(row_idx, 13).value in (None, "")
    wb.close()

    store2 = ExcelLogStore(tmp_path / "queue.json")
    row = ExcelRow(
        worksheet="SATCOM",
        collection_event="Haifa Port",
        target_class="Iridium",
        start_hz=mapping.start_hz,
        end_hz=mapping.end_hz,
        start_style=mapping.start_style,
        stop_style=mapping.stop_style,
        sensor="CRFS 100-18",
        filename="iq_task_0004.wav",
        remote_path="/mnt/1/remdata/20260915/iq_task_0004.wav",
        sensor_id="rfeyeDEMO",
        duration_s=0.08,
    )
    store2.enqueue(pending_from_row(workbook=path, row=row))

    def _busy(*_args, **_kwargs):
        raise ExcelBusyError("locked")

    monkeypatch.setattr("crfs_iq_recorder.excel_log.write_excel_rows", _busy)
    saved, pending = flush_excel_queue(store2)
    assert pending == 1
    store3 = ExcelLogStore(tmp_path / "queue.json")
    assert store3.pending_count() == 1
    assert store3.queue[0].duration_s == 0.08
    monkeypatch.setattr("crfs_iq_recorder.excel_log.write_excel_rows", append_rows_openpyxl)
    saved, pending = flush_excel_queue(store3)
    assert pending == 0
    assert saved == 1
    wb = load_workbook(path)
    iq_cells = [
        cell
        for cell in wb["SATCOM"]["I"]
        if cell.row > 1 and any(name.lower().endswith(".wav") for name in split_iq_filenames(cell.value))
    ]
    assert len(iq_cells) == 1
    assert split_iq_filenames(iq_cells[0].value) == [
        "iq_task_0001.wav",
        "iq_task_0002.wav",
        "iq_task_0003.wav",
        "iq_task_0004.wav",
    ]
    assert wb["SATCOM"].cell(iq_cells[0].row, 11).value == "0.08 seconds"
    wb.close()


def test_file_watch_finalizes_stable_parts_and_waits_for_later_ones():
    now = datetime(2026, 9, 15, 12, 0, 0)
    state = FileWatchState()
    part1 = RemoteEntry("iq_a_0001.wav", "/r/iq_a_0001.wav", False, 100, now)
    state, newly = update_file_watch(state, [part1], now=now, stable_needed=2)
    assert newly == []
    state, newly = update_file_watch(state, [part1], now=now + timedelta(seconds=2), stable_needed=2)
    assert [item.path for item in newly] == ["/r/iq_a_0001.wav"]
    assert not capture_is_settled(state, now=now + timedelta(seconds=2), elapsed=True, settle_s=8)
    part2 = RemoteEntry("iq_a_0002.wav", "/r/iq_a_0002.wav", False, 100, now + timedelta(seconds=6))
    state, newly = update_file_watch(
        state,
        [part1, part2],
        now=now + timedelta(seconds=6),
        stable_needed=2,
    )
    assert newly == []
    state, newly = update_file_watch(
        state,
        [part1, part2],
        now=now + timedelta(seconds=8),
        stable_needed=2,
    )
    assert [item.path for item in newly] == ["/r/iq_a_0002.wav"]
    assert capture_is_settled(
        state,
        now=now + timedelta(seconds=16),
        elapsed=True,
        settle_s=8,
        elapsed_at=now,
    )


@pytest.mark.skipif(not HAIFA.is_file(), reason="Haifa workbook is not on this PC")
def test_haifa_workbook_copy_preserves_vba_and_template_rows(tmp_path):
    copy_path = tmp_path / HAIFA.name
    copy2(HAIFA, copy_path)
    catalog = load_workbook_catalog(copy_path)
    classes = {(item.worksheet, item.target_class) for item in catalog.mappings}
    assert ("SATCOM", "Iridium") in classes
    assert ("SATCOM", "Satcom") in classes
    assert preferred_event_name(catalog.events) == "Haifa Port"
    original = load_workbook(copy_path, keep_vba=True)
    assert original.vba_archive is not None
    tactical_last = original["Tactical Radio"].max_row
    original.close()
    row = ExcelRow(
        worksheet="Tactical Radio",
        collection_event="Haifa Port",
        target_class="Tactical Radio",
        start_hz=157_000_000,
        end_hz=161_000_000,
        start_style=FreqTextStyle("MHz", True),
        stop_style=FreqTextStyle("MHz", True),
        sensor="CRFS 100-18",
        filename="iq_sim_0001.wav",
        remote_path="/mnt/1/remdata/20260915/iq_sim_0001.wav",
        sensor_id="sim",
        duration_s=0.4,
    )
    append_rows_openpyxl(copy_path, [row])
    updated = load_workbook(copy_path, keep_vba=True)
    assert updated.vba_archive is not None
    ws = updated["Tactical Radio"]
    assert ws["I2"].value in (None, "")
    names = [cell.value for cell in ws["I"] if cell.value]
    assert "iq_sim_0001.wav" in names
    iq_row = next(cell.row for cell in ws["I"] if cell.value == "iq_sim_0001.wav")
    assert iq_row > tactical_last or ws.cell(iq_row, 3).value == "157 MHz"
    assert ws.cell(iq_row, 3).value == "157 MHz"
    assert ws.cell(iq_row, 4).value == "161 MHz"
    assert ws.cell(iq_row, 11).value == "0.4 seconds"
    assert ws.cell(iq_row, 12).value in (None, "")
    assert ws.cell(iq_row, 13).value in (None, "")
    assert str(updated["Tactical Radio"]["M2"].value or "").startswith("=IF(K2")
    updated.close()


def preferred_event_name(events) -> str:
    from crfs_iq_recorder.collection_catalog import preferred_event

    return preferred_event(events)
