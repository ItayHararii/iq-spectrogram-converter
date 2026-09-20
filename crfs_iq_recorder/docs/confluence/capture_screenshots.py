"""Capture Confluence screenshots from demo widgets. No live sensor, IP, or password."""

from __future__ import annotations

import os
import sys
import tempfile
import time
from datetime import datetime
from pathlib import Path

os.environ.pop("QT_QPA_PLATFORM", None)

ROOT = Path(__file__).resolve().parents[2]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from PySide6.QtWidgets import QApplication

from openpyxl import Workbook

from crfs_iq_recorder.connection_state import ConnectionState
from crfs_iq_recorder.gui import MainWindow, configure_appearance
from crfs_iq_recorder.recording_history import add_recording, new_recording
from crfs_iq_recorder.sensor_info import SensorInfo
from crfs_iq_recorder.sftp_window import SftpWindow

EXAMPLE_HOST = "example"
EXAMPLE_DOWNLOAD = r"D:\CRFS IQ Recorder\Recordings"
EXAMPLE_WORKBOOK = "Collection workbook.xlsm"


def _wait(app: QApplication, pred, timeout: float = 8.0) -> bool:
    deadline = time.time() + timeout
    while time.time() < deadline:
        app.processEvents()
        if pred():
            return True
        time.sleep(0.05)
    return False


def _save(widget, path: Path) -> None:
    pix = widget.grab()
    path.parent.mkdir(parents=True, exist_ok=True)
    if not pix.save(str(path), "PNG"):
        raise SystemExit(f"Could not save {path}")
    print(f"Wrote {path} ({path.stat().st_size} bytes)")


def _sample_workbook(path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "SATCOM"
    headers = [
        "Collection Event",
        "Class",
        "Freq Start",
        "Freq Stop",
        "RBW",
        "Scan Rate",
        "Sensor",
        "Snapshot ID",
        "IQ Recording",
        "IQ Image",
        "Playback Start",
        "Playback End",
        "Playback Time (UTC-3)",
    ]
    for col, header in enumerate(headers, start=1):
        ws.cell(1, col, header)
    ws.cell(2, 1, "Haifa Port")
    ws.cell(2, 2, "Iridium")
    ws.cell(2, 3, "1616 MHz")
    ws.cell(2, 4, "1626 MHz")
    ws.cell(2, 7, "CRFS 100-18")
    wb.save(path)
    return path


def main() -> int:
    out = Path(__file__).resolve().parent
    tmp = Path(tempfile.mkdtemp(prefix="crfs-docs-shot-"))
    os.environ["QT_ENABLE_HIGHDPI_SCALING"] = "1"

    from crfs_iq_recorder import paths as app_paths
    from crfs_iq_recorder import recording_history

    app_paths.settings_path = lambda: tmp / "settings.json"  # type: ignore[method-assign]
    recording_history.recordings_path = lambda: tmp / "recordings.json"  # type: ignore[method-assign]
    add_recording(
        new_recording(
            host=EXAMPLE_HOST,
            start_hz=793_750_000,
            end_hz=806_250_000,
            center_hz=800_000_000,
            bandwidth_hz=12_500_000,
            duration_s=0.1,
            started_at=datetime(2026, 9, 10, 16, 16, 29),
            demo=True,
            iq_stem="iq_20260910_161629",
        )
    )

    app = QApplication.instance() or QApplication([])
    configure_appearance(app)

    window = MainWindow(demo=True)
    window._conn.host = EXAMPLE_HOST
    window._conn.download_dir = EXAMPLE_DOWNLOAD
    window.conn_label.setText(EXAMPLE_HOST)
    window._apply_sensor_info(
        SensorInfo(
            host=EXAMPLE_HOST,
            model="R40-8",
            firmware="2.25-325",
            serial="rfeyeDEMO",
            status="Connected",
            connected=True,
        )
    )
    window._refresh_storage()
    window.mode_start.setChecked(True)
    window.field_a.setText("1616")
    window.field_b.setText("1626")
    window.time_edit.setText("2")
    window.excel_check.setChecked(True)
    book = tmp / EXAMPLE_WORKBOOK
    _sample_workbook(book)
    window._load_workbook(str(book), quiet=True)
    window.workbook_edit.setText(r"C:\Collection\Collection workbook.xlsm")
    window._refresh_class_suggestion()
    window._refresh_size()
    window._update_collection_status()
    window.resize(1100, 1180)
    window.show()
    window.raise_()
    window.activateWindow()
    app.processEvents()
    time.sleep(0.4)
    app.processEvents()
    _save(window, out / "01-main-window.png")

    sftp = SftpWindow(
        ConnectionState(
            demo=True,
            host=EXAMPLE_HOST,
            download_dir=EXAMPLE_DOWNLOAD,
        )
    )
    sftp.resize(1240, 720)
    sftp.show()
    sftp.raise_()
    app.processEvents()
    time.sleep(0.3)
    _wait(app, lambda: sftp.table.topLevelItemCount() >= 2)
    files = [entry for entry in sftp._entries if not entry.is_dir]
    group = next(
        (
            sftp.table.topLevelItem(i)
            for i in range(sftp.table.topLevelItemCount())
            if sftp.table.topLevelItem(i) and sftp.table.topLevelItem(i).childCount()
        ),
        sftp.table.topLevelItem(0),
    )
    if group is not None:
        sftp.table.clearSelection()
        group.setExpanded(True)
        group.setSelected(True)
        sftp.table.setCurrentItem(group)
        sftp._on_selection_changed()
    _wait(
        app,
        lambda: sftp.preview_image.pixmap() is not None and not sftp.preview_image.pixmap().isNull(),
        10,
    )
    app.processEvents()
    _save(sftp, out / "02-sensor-files.png")
    _save(sftp.preview_image, out / "03-waterfall-preview.png")

    sftp.close()
    window.close()
    app.processEvents()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
