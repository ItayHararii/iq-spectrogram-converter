"""Read class/frequency mappings from a collection workbook."""

from __future__ import annotations

from collections import Counter
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable

from .frequency import FreqTextStyle, FrequencyError, freq_text_style, parse_freq_text

HEADER_COLLECTION = "Collection Event"
HEADER_CLASS = "Class"
HEADER_START = "Freq Start"
HEADER_STOP = "Freq Stop"
HEADER_RBW = "RBW"
HEADER_SCAN = "Scan Rate"
HEADER_SENSOR = "Sensor"
HEADER_SNAPSHOT = "Snapshot ID"
HEADER_IQ = "IQ Recording"
HEADER_IMAGE = "IQ Image"
HEADER_PLAY_START = "Playback Start"
HEADER_PLAY_END = "Playback End"
HEADER_PLAY_TIME = "Playback Time (UTC-3)"

REQUIRED_HEADERS = (
    HEADER_COLLECTION,
    HEADER_CLASS,
    HEADER_START,
    HEADER_STOP,
    HEADER_SENSOR,
    HEADER_IQ,
)


@dataclass(frozen=True)
class ClassMapping:
    worksheet: str
    target_class: str
    start_hz: int
    end_hz: int
    start_style: FreqTextStyle
    stop_style: FreqTextStyle


@dataclass(frozen=True)
class WorkbookCatalog:
    path: str
    mappings: tuple[ClassMapping, ...]
    events: tuple[str, ...]
    worksheets: tuple[str, ...]
    sensor_names: tuple[str, ...]

    def mappings_for_class(self, name: str) -> list[ClassMapping]:
        return [item for item in self.mappings if item.target_class == name]


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _header_map(headers: Iterable[Any]) -> dict[str, int]:
    found: dict[str, int] = {}
    for index, value in enumerate(headers):
        name = _cell_text(value)
        if name and name not in found:
            found[name] = index
    return found


def _has_required_headers(headers: dict[str, int]) -> bool:
    return all(name in headers for name in REQUIRED_HEADERS)


def parse_sheet_rows(
    worksheet: str,
    header_row: Iterable[Any],
    data_rows: Iterable[Iterable[Any]],
) -> tuple[list[ClassMapping], list[str], list[str]]:
    headers = _header_map(header_row)
    if not _has_required_headers(headers):
        return [], [], []
    mappings: list[ClassMapping] = []
    seen: set[tuple[str, str, int, int]] = set()
    events: list[str] = []
    sensors: list[str] = []
    for raw in data_rows:
        row = list(raw)
        def col(name: str) -> str:
            index = headers[name]
            return _cell_text(row[index] if index < len(row) else "")

        event = col(HEADER_COLLECTION)
        target = col(HEADER_CLASS)
        start_text = col(HEADER_START)
        stop_text = col(HEADER_STOP)
        sensor = col(HEADER_SENSOR)
        if event:
            events.append(event)
        if sensor:
            sensors.append(sensor)
        if HEADER_IQ in headers and col(HEADER_IQ):
            continue
        if not target or not start_text or not stop_text:
            continue
        try:
            start_hz = parse_freq_text(start_text, field="Freq Start")
            end_hz = parse_freq_text(stop_text, field="Freq Stop")
        except FrequencyError:
            continue
        if start_hz >= end_hz:
            continue
        start_style = freq_text_style(start_text) or FreqTextStyle("MHz", True)
        stop_style = freq_text_style(stop_text) or start_style
        key = (worksheet, target, start_hz, end_hz)
        if key in seen:
            continue
        seen.add(key)
        mappings.append(
            ClassMapping(
                worksheet=worksheet,
                target_class=target,
                start_hz=start_hz,
                end_hz=end_hz,
                start_style=start_style,
                stop_style=stop_style,
            )
        )
    return mappings, events, sensors


def _unique_keep_order(values: Iterable[str]) -> tuple[str, ...]:
    out: list[str] = []
    seen: set[str] = set()
    for item in values:
        if item in seen:
            continue
        seen.add(item)
        out.append(item)
    return tuple(out)


def preferred_event(events: Iterable[str]) -> str:
    items = [item for item in events if item]
    if not items:
        return ""
    counts = Counter(items)
    best = counts.most_common(1)[0][0]
    for item in items:
        if item == best:
            return item
    return items[0]


def load_catalog_from_rows(
    path: str,
    sheets: Iterable[tuple[str, Iterable[Any], Iterable[Iterable[Any]]]],
) -> WorkbookCatalog:
    mappings: list[ClassMapping] = []
    events: list[str] = []
    sensors: list[str] = []
    worksheets: list[str] = []
    for name, header, rows in sheets:
        parsed, sheet_events, sheet_sensors = parse_sheet_rows(name, header, rows)
        if not parsed and not sheet_events:
            continue
        worksheets.append(name)
        mappings.extend(parsed)
        events.extend(sheet_events)
        sensors.extend(sheet_sensors)
    return WorkbookCatalog(
        path=path,
        mappings=tuple(mappings),
        events=_unique_keep_order(events),
        worksheets=tuple(worksheets),
        sensor_names=_unique_keep_order(sensors),
    )


def load_workbook_catalog(path: str | Path) -> WorkbookCatalog:
    from openpyxl import load_workbook

    file = Path(path)
    workbook = load_workbook(file, read_only=True, data_only=True, keep_vba=False)
    try:
        sheets: list[tuple[str, list[Any], list[list[Any]]]] = []
        for name in workbook.sheetnames:
            ws = workbook[name]
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                continue
            sheets.append((name, list(rows[0]), [list(row) for row in rows[1:]]))
        return load_catalog_from_rows(str(file), sheets)
    finally:
        workbook.close()
