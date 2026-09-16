"""Append IQ rows to a collection workbook with openpyxl. Preserves VBA."""

from __future__ import annotations

from copy import copy
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .collection_catalog import (
    HEADER_CLASS,
    HEADER_COLLECTION,
    HEADER_IQ,
    HEADER_PLAY_END,
    HEADER_PLAY_START,
    HEADER_PLAY_TIME,
    HEADER_SENSOR,
    HEADER_START,
    HEADER_STOP,
    REQUIRED_HEADERS,
)
from .frequency import FreqTextStyle, format_freq_text
from .recording_history import join_iq_filenames, sort_iq_filenames, split_iq_filenames


class ExcelBusyError(Exception):
    """Workbook is locked, Excel is busy, or the file cannot be saved now."""


@dataclass(frozen=True)
class ExcelRow:
    worksheet: str
    collection_event: str
    target_class: str
    start_hz: int
    end_hz: int
    start_style: FreqTextStyle
    stop_style: FreqTextStyle
    sensor: str
    filename: str
    local_path: str = ""
    remote_path: str = ""
    sensor_id: str = ""
    duration_s: float | None = None
    record_id: str = ""
    group_key: str = ""


def backup_workbook(path: Path, dest: Path) -> Path:
    dest.parent.mkdir(parents=True, exist_ok=True)
    from shutil import copy2

    copy2(path, dest)
    return dest


def _header_map(ws: Worksheet) -> dict[str, int]:
    found: dict[str, int] = {}
    for col in range(1, (ws.max_column or 1) + 1):
        name = ws.cell(1, col).value
        text = str(name).strip() if name is not None else ""
        if text and text not in found:
            found[text] = col
    return found


def _last_used_row(ws: Worksheet, columns: int) -> int:
    last = 1
    for row in range(1, (ws.max_row or 1) + 1):
        for col in range(1, columns + 1):
            if ws.cell(row, col).value not in (None, ""):
                last = row
                break
    return last


def existing_iq_names(ws: Worksheet, headers: dict[str, int]) -> set[str]:
    iq_col = headers.get(HEADER_IQ)
    if not iq_col:
        return set()
    names: set[str] = set()
    for row in range(2, (ws.max_row or 1) + 1):
        for name in split_iq_filenames(ws.cell(row, iq_col).value):
            names.add(name)
    return names


def _copy_row_style(ws: Worksheet, src_row: int, dest_row: int, columns: int) -> None:
    src_dim = ws.row_dimensions[src_row]
    dest_dim = ws.row_dimensions[dest_row]
    if src_dim.height:
        dest_dim.height = src_dim.height
    for col in range(1, columns + 1):
        src = ws.cell(src_row, col)
        dest = ws.cell(dest_row, col)
        if src.has_style:
            dest.font = copy(src.font)
            dest.border = copy(src.border)
            dest.fill = copy(src.fill)
            dest.number_format = src.number_format
            dest.protection = copy(src.protection)
            dest.alignment = copy(src.alignment)


def _set_iq_cell(cell, filename: str, local_path: str) -> None:
    names = split_iq_filenames(filename)
    cell.value = join_iq_filenames(names) if names else filename
    align = copy(cell.alignment) if cell.has_style else Alignment()
    align.wrap_text = True
    align.vertical = "center"
    cell.alignment = align
    local = Path(local_path) if local_path else None
    if len(names) <= 1 and local is not None and local.is_file():
        cell.hyperlink = str(local.resolve())
    elif cell.hyperlink:
        cell.hyperlink = None


def _fit_iq_row(ws: Worksheet, row: int, filename: str) -> None:
    lines = max(1, len(split_iq_filenames(filename) or [filename]))
    ws.row_dimensions[row].height = max(18, 16 * lines + 4)


def format_playback_duration(duration_s: float | None) -> str | None:
    """Plain text for Playback Start, for example '2 seconds'."""
    if duration_s is None:
        return None
    try:
        value = float(duration_s)
    except (TypeError, ValueError):
        return None
    if value != value or value < 0:
        return None
    if abs(value - round(value)) < 1e-6:
        number = str(int(round(value)))
    else:
        number = f"{value:g}"
    return f"{number} seconds"


def _set_iq_duration_cell(cell, duration_s: float | None) -> None:
    text = format_playback_duration(duration_s)
    if text is None:
        return
    cell.value = text
    cell.number_format = "General"


def find_row_with_iq_names(ws: Worksheet, names: list[str]) -> int | None:
    headers = _header_map(ws)
    iq_col = headers.get(HEADER_IQ)
    if not iq_col:
        return None
    wanted = {name.casefold() for name in names if name}
    if not wanted:
        return None
    for row in range(2, (ws.max_row or 1) + 1):
        existing = split_iq_filenames(ws.cell(row, iq_col).value)
        if any(name.casefold() in wanted for name in existing):
            return row
    return None


def append_iq_row(ws: Worksheet, row_data: ExcelRow) -> int:
    headers = _header_map(ws)
    missing = [name for name in REQUIRED_HEADERS if name not in headers]
    if missing:
        raise ValueError(f"Worksheet {ws.title!r} is missing columns: {', '.join(missing)}")
    columns = max(headers.values())
    incoming = split_iq_filenames(row_data.filename) or [row_data.filename]
    found = find_row_with_iq_names(ws, incoming)
    if found:
        dest = found
        existing = split_iq_filenames(ws.cell(dest, headers[HEADER_IQ]).value)
        merged = join_iq_filenames(existing + incoming)
    else:
        last = _last_used_row(ws, columns)
        dest = last + 1
        style_from = last if last >= 2 else 1
        _copy_row_style(ws, style_from, dest, columns)
        merged = join_iq_filenames(incoming)

        def write(header: str, value: Any) -> None:
            if header not in headers:
                return
            ws.cell(dest, headers[header]).value = value

        write(HEADER_COLLECTION, row_data.collection_event)
        write(HEADER_CLASS, row_data.target_class)
        write(HEADER_START, format_freq_text(row_data.start_hz, row_data.start_style))
        write(HEADER_STOP, format_freq_text(row_data.end_hz, row_data.stop_style))
        write(HEADER_SENSOR, row_data.sensor)
        if HEADER_PLAY_END in headers:
            ws.cell(dest, headers[HEADER_PLAY_END]).value = None
        if HEADER_PLAY_TIME in headers:
            ws.cell(dest, headers[HEADER_PLAY_TIME]).value = None
    _set_iq_cell(ws.cell(dest, headers[HEADER_IQ]), merged, row_data.local_path)
    _fit_iq_row(ws, dest, merged)
    if HEADER_PLAY_START in headers:
        _set_iq_duration_cell(ws.cell(dest, headers[HEADER_PLAY_START]), row_data.duration_s)
    ws.auto_filter.ref = f"A1:{get_column_letter(columns)}{dest if dest > 1 else _last_used_row(ws, columns)}"
    return dest


def workbook_has_iq_filename(path: Path, sheet: str, filename: str) -> bool:
    workbook = load_workbook(path, read_only=True, data_only=True, keep_vba=False)
    try:
        if sheet not in workbook.sheetnames:
            return False
        ws = workbook[sheet]
        headers = {}
        first = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), None)
        if not first:
            return False
        for index, value in enumerate(first, start=1):
            text = str(value).strip() if value is not None else ""
            if text:
                headers[text] = index
        iq_col = headers.get(HEADER_IQ)
        if not iq_col:
            return False
        wanted = {name.casefold() for name in (split_iq_filenames(filename) or [filename]) if name}
        if not wanted:
            return False
        for row in ws.iter_rows(min_row=2, min_col=iq_col, max_col=iq_col, values_only=True):
            names = split_iq_filenames(row[0] if row else None)
            if any(name.casefold() in wanted for name in names):
                return True
        return False
    finally:
        workbook.close()


def append_rows_openpyxl(path: Path, rows: list[ExcelRow]) -> list[int]:
    if not rows:
        return []
    try:
        workbook = load_workbook(path, keep_vba=True)
    except PermissionError as exc:
        raise ExcelBusyError(f"Could not open {path.name}. It may be open in Excel.") from exc
    try:
        written: list[int] = []
        for row in rows:
            if row.worksheet not in workbook.sheetnames:
                raise ValueError(f"Workbook has no worksheet named {row.worksheet!r}.")
            written.append(append_iq_row(workbook[row.worksheet], row))
        try:
            workbook.save(path)
        except PermissionError as exc:
            raise ExcelBusyError(f"Could not save {path.name}. Excel may have it open.") from exc
        return written
    finally:
        vba = getattr(workbook, "vba_archive", None)
        if vba is not None:
            try:
                vba.close()
            except Exception:
                pass
            workbook.vba_archive = None
        workbook.close()
